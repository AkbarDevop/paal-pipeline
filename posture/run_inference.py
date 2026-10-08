#!/usr/bin/env python3
"""One-command posture inference pipeline.

Takes a single input path to raw OAK-D data, crops images, runs the
trained model, and produces a predictions CSV + posture heatmap.

Usage:
    python run_inference.py C:\\PAAL_Data\\fed_pig
    python run_inference.py /path/to/oak_data
"""

import argparse
import csv
import os
import re
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime

import numpy as np
import torch

try:
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    import matplotlib.dates as mdates
    HAS_PLT = True
except ImportError:
    HAS_PLT = False

import sys, pathlib
sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from torch.utils.data import DataLoader, Dataset

from config import CROP_TOF, MODEL_DIR, OUTPUT_DIR, POSTURE3_CLASSES, IMG_SIZE
from models import SingleModalModel
from preprocess import (ckpt_preprocessing, crop_stall, load_model_input, normalize_batch,
                        pick_device, read_image, resize_rgb_uint8)
from id_corrections import assign_ids, dedupe, folder_quality
from temporal import smooth_results

# ── Constants ─────────────────────────────────────────────────────────────

TOF_W, TOF_H = 640, 480
DEPTH_THRESHOLD = 1463
BAR_MARGIN = 50

FILENAME_RE = re.compile(
    r"^pig(\d+)_(depth_vis|depth|ir_vis|ir|rgb_aligned|rgb)_"
    r"(\d{8}-\d{2}-\d{2}-\d{2})(_cropped)?\.(jpg|raw)$"
)
FOLDER_RE = re.compile(r"^\d{8}-\d{2}-\d{2}-\d{2}$")


# ── Step 1: Scan + load (crop happens in memory, raw folders are read-only) ──

def load_and_crop_ir(ir_path, size=IMG_SIZE):
    """Kept for backward compatibility (Colab / older scripts).

    Returns an HWC float32 array in 0..1, exactly like v1.0.
    """
    t = load_model_input(ir_path, size=size, crop=True, normalize=None)
    return None if t is None else t.permute(1, 2, 0).numpy()


def load_depth_raw(path):
    if not path or not os.path.exists(path):
        return None
    size = os.path.getsize(path)
    expected = TOF_W * TOF_H * 2
    if size == expected + 8:
        raw = np.fromfile(path, dtype=np.uint16, offset=8)
    elif size == expected:
        raw = np.fromfile(path, dtype=np.uint16)
    else:
        return None
    if raw.size != TOF_W * TOF_H:
        return None
    return raw.reshape((TOF_H, TOF_W))


def check_pig_present(depth_raw_path):
    """Median-depth prefilter. Missing/unreadable depth -> assume present."""
    depth = load_depth_raw(depth_raw_path)
    if depth is None:
        return True, 0.0
    x1, y1, x2, y2 = CROP_TOF
    crop = depth[y1:y2, x1:x2].copy()
    crop[:, :BAR_MARGIN] = 0
    crop[:, -BAR_MARGIN:] = 0
    valid = crop[(crop > 200) & (crop < 5000)]
    if len(valid) == 0:
        return False, 0.0
    median = float(np.median(valid))
    return median < DEPTH_THRESHOLD, median


def scan_folder(data_dir):
    folders = sorted(
        d for d in os.listdir(data_dir)
        if os.path.isdir(os.path.join(data_dir, d)) and FOLDER_RE.match(d)
    )
    if not folders:
        return []

    records = {}
    for folder in folders:
        folder_path = os.path.join(data_dir, folder)
        for fname in os.listdir(folder_path):
            m = FILENAME_RE.match(fname)
            if not m:
                continue
            pig_id = int(m.group(1))
            modality = m.group(2)
            timestamp = m.group(3)
            is_cropped = m.group(4) is not None
            ext = m.group(5)

            key = (folder, pig_id, timestamp)
            if key not in records:
                records[key] = {
                    "timestamp_folder": folder,
                    "pig_id": pig_id,
                    "pig_timestamp": timestamp,
                }

            if modality == "ir_vis":
                mod_key = "ir"
            elif modality == "depth_vis":
                mod_key = "depth"
            else:
                mod_key = modality

            col_name = f"{mod_key}{'_cropped' if is_cropped else ''}_{ext}"
            records[key][col_name] = os.path.join(folder_path, fname)

    return sorted(records.values(), key=lambda r: (r["timestamp_folder"], r["pig_id"], r["pig_timestamp"]))


def get_image_path(frame):
    """Uncropped IR path, falling back to a legacy _cropped file."""
    for key in ("ir_jpg", "ir_cropped_jpg"):
        p = frame.get(key, "")
        if p and os.path.exists(p):
            return p
    return ""


def load_pig_detector(model_dir, device):
    """Load the binary pig-presence CNN. Returns None if weights not found."""
    path = os.path.join(model_dir, "pig_detector.pth")
    if not os.path.exists(path):
        return None
    import torchvision.models as tv_models
    detector = tv_models.mobilenet_v2(weights=None)
    detector.classifier[1] = torch.nn.Linear(detector.classifier[1].in_features, 2)
    ckpt = torch.load(path, map_location=device, weights_only=False)
    detector.load_state_dict(ckpt["model_state_dict"])
    detector.eval().to(device)
    detector.normalize = ckpt.get("normalize", None)
    return detector


# status codes returned by FrameDataset
ST_OK, ST_EMPTY_DEPTH, ST_CORRUPT = 0, 1, 2


class FrameDataset(Dataset):
    """Loads one capture: depth prefilter + IR as uint8 CHW.

    Returns the stall-cropped image (what the pig detector was trained on) and
    the image for the posture model (cropped unless its checkpoint says
    crop=False). Lives at module level so DataLoader workers can pickle it on
    Windows.
    """

    def __init__(self, frames, size=IMG_SIZE, crop_for_model=True):
        self.frames = frames
        self.size = size
        self.crop_for_model = crop_for_model

    def __len__(self):
        return len(self.frames)

    def __getitem__(self, idx):
        frame = self.frames[idx]
        empty = torch.zeros((3, self.size, self.size), dtype=torch.uint8)
        present, median_d = check_pig_present(frame.get("depth_raw", ""))
        if not present:
            return idx, empty, empty, ST_EMPTY_DEPTH, median_d
        path = get_image_path(frame)
        img = read_image(path)
        if img is None:
            return idx, empty, empty, ST_CORRUPT, median_d

        def to_t(a):
            return torch.from_numpy(np.ascontiguousarray(a.transpose(2, 0, 1)))

        cropped = to_t(resize_rgb_uint8(crop_stall(img, path), self.size))
        model_img = cropped if self.crop_for_model else to_t(resize_rgb_uint8(img, self.size))
        return idx, cropped, model_img, ST_OK, median_d


def run_predictions(data_dir, device, model, detector=None, batch_size=64, workers=4,
                    posture_normalize=None, frames=None, posture_crop=True):
    """Batched inference. Returns (results, stats).

    Same decisions as v1.0 (depth prefilter -> CNN detector p>=0.5 -> posture),
    just batched and with full class probabilities kept.
    """
    if frames is None:
        frames = scan_folder(data_dir)
    if not frames:
        print("No frames found.")
        return [], {}

    print(f"[1/2] Running inference on {len(frames)} frames "
          f"(batch {batch_size}, {workers} loader workers, {device})...")
    if detector is None:
        print("  (pig presence: depth prefilter only — detector model not loaded)")
    else:
        print("  (pig presence: depth prefilter + CNN detector)")

    loader = DataLoader(FrameDataset(frames, crop_for_model=posture_crop), batch_size=batch_size, shuffle=False,
                        num_workers=workers, persistent_workers=False)
    class_ids = sorted(POSTURE3_CLASSES)
    results = []
    stats = {"scanned": len(frames), "empty_depth": 0, "empty_detector": 0, "corrupt": 0}
    done = 0
    t0 = time.time()

    with torch.inference_mode():
        for idxs, imgs, model_imgs, status, medians in loader:
            done += len(idxs)
            stats["empty_depth"] += int((status == ST_EMPTY_DEPTH).sum())
            stats["corrupt"] += int((status == ST_CORRUPT).sum())
            keep = (status == ST_OK).nonzero(as_tuple=True)[0]
            if len(keep) == 0:
                continue
            x = imgs[keep].to(device).float().div_(255.0)
            xm = x if posture_crop else model_imgs[keep].to(device).float().div_(255.0)

            present_prob = torch.ones(len(keep))
            if detector is not None:
                det = torch.softmax(detector(normalize_batch(x, detector.normalize)), dim=1)[:, 1].cpu()
                present_prob = det
                ok = det >= 0.5
                stats["empty_detector"] += int((~ok).sum())
                keep, xm, present_prob = keep[ok], xm[ok.to(xm.device)], det[ok]
                if len(keep) == 0:
                    continue

            probs = torch.softmax(model(normalize_batch(xm, posture_normalize)), dim=1).cpu().numpy()
            for j, k in enumerate(keep.tolist()):
                frame = frames[int(idxs[k])]
                pred = int(np.argmax(probs[j]))
                row = {
                    "timestamp_folder": frame["timestamp_folder"],
                    "pig_id": frame["pig_id"],
                    "pig_timestamp": frame["pig_timestamp"],
                    "prediction": pred,
                    "prediction_name": POSTURE3_CLASSES.get(pred, str(pred)),
                    "confidence": round(float(probs[j][pred]), 4),
                    "median_depth": round(float(medians[k]), 1),
                    "image_path": get_image_path(frame),
                    "pig_present_prob": round(float(present_prob[j]), 4),
                }
                for c in class_ids:
                    row[f"prob_{POSTURE3_CLASSES[c]}"] = round(float(probs[j][c]), 4)
                results.append(row)

            if done % (batch_size * 20) < batch_size:
                rate = done / max(time.time() - t0, 1e-6)
                print(f"  {done}/{len(frames)} frames ({rate:.0f} frames/s)")

    results.sort(key=lambda r: (r["timestamp_folder"], r["pig_id"], r["pig_timestamp"]))
    print(f"  Predicted: {len(results)}, "
          f"Skipped (depth prefilter): {stats['empty_depth']}, "
          f"Skipped (CNN detector): {stats['empty_detector']}, "
          f"Skipped (corrupt/missing): {stats['corrupt']} "
          f"[{time.time() - t0:.0f}s]")
    return results, stats


# ── Step 2b: ID policy + per-pig sequence columns ──────────────────────────

def add_change_column(results, key="prediction_name", out_key="posture_changed", max_gap_minutes=30):
    """1 if the posture differs from the same pig's previous kept frame (within max_gap)."""
    by_pig = defaultdict(list)
    for r in results:
        by_pig[r["pig_id"]].append(r)
    for rows in by_pig.values():
        rows.sort(key=lambda r: r["pig_timestamp"])
        prev = None
        for r in rows:
            t = parse_timestamp(r["pig_timestamp"])
            changed = 0
            if prev is not None and t is not None and prev[1] is not None:
                gap = (t - prev[1]).total_seconds() / 60.0
                if gap <= max_gap_minutes and r[key] != prev[0]:
                    changed = 1
            r[out_key] = changed
            prev = (r[key], t)
    return results


def save_rows(rows, out_path, fields):
    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        writer.writeheader()
        writer.writerows(rows)


def save_coverage(results, out_path, num_stalls=20):
    """Kept frames per stall per calendar day — shows days a stall has no data."""
    counts = defaultdict(lambda: defaultdict(int))
    days = set()
    for r in results:
        day = r["pig_timestamp"][:8]
        days.add(day)
        counts[r["pig_id"]][day] += 1
    days = sorted(days)
    with open(out_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["pig_id"] + days)
        for pid in range(num_stalls):
            w.writerow([pid] + [counts[pid].get(d, 0) for d in days])


def save_posture_matrix(results, out_path, num_stalls=20, key="prediction_name"):
    """Wide table Fengkai asked for: one row per timestamp folder, columns pig0..pig19."""
    grid = defaultdict(dict)
    for r in results:
        grid[r["timestamp_folder"]][r["pig_id"]] = r[key]
    with open(out_path, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(["timestamp_folder", "datetime"] + [f"pig{i}" for i in range(num_stalls)])
        for folder in sorted(grid):
            t = parse_timestamp(folder)
            w.writerow([folder, t.strftime("%Y-%m-%d %H:%M:%S") if t else ""]
                       + [grid[folder].get(i, "") for i in range(num_stalls)])


PRED_FIELDS = ["timestamp_folder", "pig_id", "pig_timestamp",
               "prediction", "prediction_name", "confidence",
               "median_depth", "image_path",
               "pig_id_raw", "id_status", "pig_present_prob",
               "prob_standing", "prob_sitting", "prob_lying",
               "posture_changed"]
SMOOTH_FIELDS = ["smoothed_prediction_name", "posture_changed_smoothed"]


# ── Step 3: Generate outputs ──────────────────────────────────────────────

def save_csv(results, out_path):
    save_rows(results, out_path, PRED_FIELDS)


def parse_timestamp(ts_str):
    """Parse pig_timestamp like '20260211-09-17-49' to datetime."""
    try:
        return datetime.strptime(ts_str, "%Y%m%d-%H-%M-%S")
    except ValueError:
        return None


def normalize_pig_id(pid, num_stalls=20):
    """Camera counter can overflow past num_stalls. Map back with modulo."""
    return pid % num_stalls


def generate_heatmap(results, out_path):
    """Generate a posture timeline heatmap (time x pig_id)."""
    if not HAS_PLT:
        print("  matplotlib not installed, skipping heatmap")
        return

    # IDs are already corrected by id_corrections.assign_ids (no blind % 20)
    # Fixed Y-axis: always pig 0–19 regardless of which pigs have data
    pig_ids = list(range(20))
    timestamps = [parse_timestamp(r["pig_timestamp"]) for r in results]
    valid = [(r, t) for r, t in zip(results, timestamps) if t is not None]
    if not valid:
        print("  No valid timestamps for heatmap")
        return

    min_t = min(t for _, t in valid)
    max_t = max(t for _, t in valid)

    # Create hour-level bins
    from datetime import timedelta
    hours = []
    t = min_t.replace(minute=0, second=0)
    while t <= max_t:
        hours.append(t)
        t += timedelta(hours=1)

    if not hours or not pig_ids:
        return

    # Posture encoding: standing=0, sitting=1, lying=2, no_data=-1
    posture_map = {"standing": 0, "sitting": 1, "lying": 2}

    # Majority vote per (pig, hour) cell
    from collections import defaultdict
    votes = defaultdict(list)
    for r, t in valid:
        if t is None:
            continue
        hi = int((t - hours[0]).total_seconds() / 3600)
        hi = min(hi, len(hours) - 1)
        pid = r["pig_id"]
        posture = posture_map.get(r["prediction_name"], -1)
        votes[(pid, hi)].append(posture)

    pig_idx = {pid: i for i, pid in enumerate(pig_ids)}
    grid = np.full((len(pig_ids), len(hours)), -1, dtype=float)
    for (pid, hi), postures in votes.items():
        # Majority vote: most common posture in this hour
        grid[pig_idx[pid], hi] = max(set(postures), key=postures.count)

    # Custom colormap: gray=no_data, green=standing, orange=sitting, blue=lying
    from matplotlib.colors import ListedColormap, BoundaryNorm
    colors = ["#d4d4d4", "#22c55e", "#f97316", "#3b82f6"]
    cmap = ListedColormap(colors)
    bounds = [-1.5, -0.5, 0.5, 1.5, 2.5]
    norm = BoundaryNorm(bounds, cmap.N)

    fig, ax = plt.subplots(figsize=(max(14, len(hours) * 0.3), 8))
    im = ax.pcolormesh(grid, cmap=cmap, norm=norm, edgecolors="white", linewidth=0.5)

    # Y-axis: fixed pig 0–19
    ax.set_yticks(np.arange(len(pig_ids)) + 0.5)
    ax.set_yticklabels([f"pig {pid}" for pid in pig_ids], fontsize=8)
    ax.set_ylim(0, len(pig_ids))

    # X-axis: fixed 6-hour tick intervals
    step = max(1, 6)
    ax.set_xticks(np.arange(0, len(hours), step) + 0.5)
    ax.set_xticklabels(
        [hours[i].strftime("%m/%d %H:%M") for i in range(0, len(hours), step)],
        rotation=45, ha="right", fontsize=7,
    )
    ax.set_xlim(0, len(hours))

    ax.set_xlabel("Time")
    ax.set_ylabel("Pig ID")
    ax.set_title("Posture Timeline Heatmap (hourly majority vote)")

    # Legend
    from matplotlib.patches import Patch
    legend_elements = [
        Patch(facecolor="#22c55e", label="Standing"),
        Patch(facecolor="#f97316", label="Sitting"),
        Patch(facecolor="#3b82f6", label="Lying"),
        Patch(facecolor="#d4d4d4", label="No data"),
    ]
    ax.legend(handles=legend_elements, loc="upper right", fontsize=8)

    plt.tight_layout()
    plt.savefig(out_path, dpi=150)
    plt.close()


def generate_heatmap_raw(results, out_path):
    """Generate heatmap with RAW pig IDs (including overflow 20+).

    This shows all pig IDs the camera assigned, without normalization,
    so you can visually spot which timestamps have duplicates (pig20, pig21, etc).
    """
    if not HAS_PLT:
        print("  matplotlib not installed, skipping raw heatmap")
        return

    from datetime import timedelta
    from collections import defaultdict
    from matplotlib.colors import ListedColormap, BoundaryNorm
    from matplotlib.patches import Patch

    posture_map = {"standing": 0, "sitting": 1, "lying": 2}

    # Use ALL pig IDs found in data (including overflow)
    all_pig_ids = sorted(set(r["pig_id"] for r in results))
    max_pid = max(all_pig_ids) if all_pig_ids else 19

    # Y-axis: pig 0 through max pig ID
    pig_ids = list(range(max_pid + 1))

    timestamps = [parse_timestamp(r["pig_timestamp"]) for r in results]
    valid = [(r, t) for r, t in zip(results, timestamps) if t is not None]
    if not valid:
        return

    min_t = min(t for _, t in valid)
    max_t = max(t for _, t in valid)
    hours = []
    t = min_t.replace(minute=0, second=0)
    while t <= max_t:
        hours.append(t)
        t += timedelta(hours=1)

    if not hours or not pig_ids:
        return

    votes = defaultdict(list)
    for r, t in valid:
        hi = int((t - hours[0]).total_seconds() / 3600)
        hi = min(hi, len(hours) - 1)
        pid = r["pig_id"]
        posture = posture_map.get(r["prediction_name"], -1)
        votes[(pid, hi)].append(posture)

    pig_idx = {pid: i for i, pid in enumerate(pig_ids)}
    grid = np.full((len(pig_ids), len(hours)), -1, dtype=float)
    for (pid, hi), postures in votes.items():
        if pid in pig_idx:
            grid[pig_idx[pid], hi] = max(set(postures), key=postures.count)

    colors = ["#d4d4d4", "#22c55e", "#f97316", "#3b82f6"]
    cmap = ListedColormap(colors)
    bounds = [-1.5, -0.5, 0.5, 1.5, 2.5]
    norm = BoundaryNorm(bounds, cmap.N)

    fig, ax = plt.subplots(figsize=(max(14, len(hours) * 0.3), max(8, len(pig_ids) * 0.4)))
    ax.pcolormesh(grid, cmap=cmap, norm=norm, edgecolors="white", linewidth=0.5)

    ax.set_yticks(np.arange(len(pig_ids)) + 0.5)
    labels = []
    for pid in pig_ids:
        if pid >= 20:
            labels.append(f"pig {pid} (overflow)")
        else:
            labels.append(f"pig {pid}")
    ax.set_yticklabels(labels, fontsize=7)
    ax.set_ylim(0, len(pig_ids))

    # Draw a red line at pig 20 to separate normal from overflow
    if max_pid >= 20:
        ax.axhline(y=20, color="red", linewidth=2, linestyle="--")

    step = max(1, 6)
    ax.set_xticks(np.arange(0, len(hours), step) + 0.5)
    ax.set_xticklabels(
        [hours[i].strftime("%m/%d %H:%M") for i in range(0, len(hours), step)],
        rotation=45, ha="right", fontsize=7,
    )
    ax.set_xlim(0, len(hours))

    ax.set_xlabel("Time")
    ax.set_ylabel("Camera Pig ID (raw, unnormalized)")
    ax.set_title("Posture Heatmap — RAW Pig IDs (overflow visible above red line)")

    legend_elements = [
        Patch(facecolor="#22c55e", label="Standing"),
        Patch(facecolor="#f97316", label="Sitting"),
        Patch(facecolor="#3b82f6", label="Lying"),
        Patch(facecolor="#d4d4d4", label="No data"),
    ]
    ax.legend(handles=legend_elements, loc="upper right", fontsize=8)

    plt.tight_layout()
    plt.savefig(out_path, dpi=150)
    plt.close()
    print(f"  Raw ID heatmap: {out_path}")


def generate_per_pig_heatmaps(results, out_dir):
    """Generate individual heatmap per pig ID."""
    if not HAS_PLT:
        return

    from datetime import timedelta
    from matplotlib.colors import ListedColormap, BoundaryNorm
    from matplotlib.patches import Patch

    posture_map = {"standing": 0, "sitting": 1, "lying": 2}
    colors = ["#d4d4d4", "#22c55e", "#f97316", "#3b82f6"]
    cmap = ListedColormap(colors)
    bounds = [-1.5, -0.5, 0.5, 1.5, 2.5]
    norm = BoundaryNorm(bounds, cmap.N)

    pig_ids = sorted(set(r["pig_id"] for r in results))
    timestamps = [(r, parse_timestamp(r["pig_timestamp"])) for r in results]
    valid = [(r, t) for r, t in timestamps if t is not None]
    if not valid:
        return

    min_t = min(t for _, t in valid)
    max_t = max(t for _, t in valid)
    hours = []
    t = min_t.replace(minute=0, second=0)
    while t <= max_t:
        hours.append(t)
        t += timedelta(hours=1)

    pig_dir = os.path.join(out_dir, "per_pig")
    os.makedirs(pig_dir, exist_ok=True)

    for pid in pig_ids:
        pig_results = [(r, t) for r, t in valid if r["pig_id"] == pid]
        if not pig_results:
            continue

        from collections import defaultdict
        votes = defaultdict(list)
        for r, t in pig_results:
            hi = int((t - hours[0]).total_seconds() / 3600)
            hi = min(hi, len(hours) - 1)
            votes[hi].append(posture_map.get(r["prediction_name"], -1))

        grid = np.full((1, len(hours)), -1, dtype=float)
        for hi, postures in votes.items():
            grid[0, hi] = max(set(postures), key=postures.count)

        _, ax = plt.subplots(figsize=(max(14, len(hours) * 0.15), 2))
        ax.pcolormesh(grid, cmap=cmap, norm=norm, edgecolors="white", linewidth=0.5)
        step = max(1, 6)
        ax.set_xticks(np.arange(0, len(hours), step) + 0.5)
        ax.set_xticklabels(
            [hours[i].strftime("%m/%d %H:%M") for i in range(0, len(hours), step)],
            rotation=45, ha="right", fontsize=7,
        )
        ax.set_xlim(0, len(hours))
        ax.set_yticks([0.5])
        ax.set_yticklabels([f"pig {pid}"], fontsize=9)
        ax.set_title(f"Pig {pid} — Posture Timeline", fontsize=11)

        legend_elements = [
            Patch(facecolor="#22c55e", label="Standing"),
            Patch(facecolor="#f97316", label="Sitting"),
            Patch(facecolor="#3b82f6", label="Lying"),
            Patch(facecolor="#d4d4d4", label="No data"),
        ]
        ax.legend(handles=legend_elements, loc="upper right", fontsize=7, ncol=4)
        plt.tight_layout()
        plt.savefig(os.path.join(pig_dir, f"pig{pid}_heatmap.png"), dpi=150)
        plt.close()

    print(f"  Per-pig heatmaps: {pig_dir}/ ({len(pig_ids)} pigs)")


def generate_per_pig_csvs(results, out_dir):
    """Generate a separate CSV file for each pig ID."""
    pig_dir = os.path.join(out_dir, "per_pig")
    os.makedirs(pig_dir, exist_ok=True)

    pig_ids = sorted(set(r["pig_id"] for r in results))
    for pid in pig_ids:
        pig_results = sorted((r for r in results if r["pig_id"] == pid),
                             key=lambda r: r["pig_timestamp"])
        csv_path = os.path.join(pig_dir, f"pig{pid}_predictions.csv")
        save_rows(pig_results, csv_path, PRED_FIELDS)

    print(f"  Per-pig CSVs: {pig_dir}/ ({len(pig_ids)} pigs)")


def generate_per_pig_excels(results, out_dir):
    """Generate a separate Excel file for each pig ID."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("  openpyxl not installed, skipping per-pig Excel files")
        return

    pig_dir = os.path.join(out_dir, "per_pig")
    os.makedirs(pig_dir, exist_ok=True)

    fill_standing = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    fill_sitting  = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
    fill_lying    = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")
    fill_header   = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    font_header   = Font(bold=True, color="FFFFFF", size=11)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    posture_fills = {"standing": fill_standing, "sitting": fill_sitting, "lying": fill_lying}

    pig_ids = sorted(set(r["pig_id"] for r in results))

    for pid in pig_ids:
        pig_results = sorted(
            [r for r in results if r["pig_id"] == pid],
            key=lambda r: r["pig_timestamp"],
        )
        if not pig_results:
            continue

        wb = Workbook()
        ws = wb.active
        ws.title = f"Pig {pid}"

        headers = ["Date", "Time", "Posture", "Confidence %", "Median Depth (mm)"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

        for i, r in enumerate(pig_results, start=2):
            t = parse_timestamp(r["pig_timestamp"])
            ws.cell(row=i, column=1, value=t.strftime("%Y-%m-%d") if t else "")
            ws.cell(row=i, column=2, value=t.strftime("%H:%M:%S") if t else "")
            posture_cell = ws.cell(row=i, column=3, value=r["prediction_name"].capitalize())
            ws.cell(row=i, column=4, value=round(r["confidence"] * 100, 1))
            ws.cell(row=i, column=5, value=r["median_depth"])

            fill = posture_fills.get(r["prediction_name"])
            if fill:
                posture_cell.fill = fill

            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).border = thin_border
                ws.cell(row=i, column=col).alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 20)

        wb.save(os.path.join(pig_dir, f"pig{pid}_report.xlsx"))

    print(f"  Per-pig Excel files: {pig_dir}/ ({len(pig_ids)} pigs)")


def generate_excel_report(results, out_path):
    """Generate an Excel workbook with summary + per-pig sheets."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        print("  openpyxl not installed, skipping Excel report (pip install openpyxl)")
        return

    from collections import defaultdict

    wb = Workbook()

    # ── Color fills for postures ──
    fill_standing = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # green
    fill_sitting  = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")  # orange
    fill_lying    = PatternFill(start_color="BDD7EE", end_color="BDD7EE", fill_type="solid")  # blue
    fill_header   = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")  # dark blue
    font_header   = Font(bold=True, color="FFFFFF", size=11)
    font_bold     = Font(bold=True, size=11)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )
    posture_fills = {"standing": fill_standing, "sitting": fill_sitting, "lying": fill_lying}

    def style_header(ws, row, ncols):
        for col in range(1, ncols + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = fill_header
            cell.font = font_header
            cell.alignment = Alignment(horizontal="center")
            cell.border = thin_border

    # ── Sheet 1: Summary (daily stats per pig) ──
    ws = wb.active
    ws.title = "Summary"

    # Build daily stats: {pig_id: {date_str: {standing, sitting, lying, total}}}
    pig_ids = sorted(set(r["pig_id"] for r in results))
    daily = defaultdict(lambda: defaultdict(lambda: {"standing": 0, "sitting": 0, "lying": 0, "total": 0}))
    all_dates = set()
    for r in results:
        t = parse_timestamp(r["pig_timestamp"])
        if t is None:
            continue
        date_str = t.strftime("%Y-%m-%d")
        all_dates.add(date_str)
        daily[r["pig_id"]][date_str][r["prediction_name"]] += 1
        daily[r["pig_id"]][date_str]["total"] += 1
    all_dates = sorted(all_dates)

    # Header
    headers = ["Pig ID", "Date", "Total Obs", "Standing", "Sitting", "Lying",
               "Standing %", "Sitting %", "Lying %"]
    for col, h in enumerate(headers, 1):
        ws.cell(row=1, column=col, value=h)
    style_header(ws, 1, len(headers))

    row = 2
    for pid in pig_ids:
        for date_str in all_dates:
            s = daily[pid][date_str]
            if s["total"] == 0:
                continue
            t = s["total"]
            ws.cell(row=row, column=1, value=f"Pig {pid}").font = font_bold
            ws.cell(row=row, column=2, value=date_str)
            ws.cell(row=row, column=3, value=t)
            ws.cell(row=row, column=4, value=s["standing"])
            ws.cell(row=row, column=5, value=s["sitting"])
            ws.cell(row=row, column=6, value=s["lying"])
            ws.cell(row=row, column=7, value=round(s["standing"] / t * 100, 1))
            ws.cell(row=row, column=8, value=round(s["sitting"] / t * 100, 1))
            ws.cell(row=row, column=9, value=round(s["lying"] / t * 100, 1))
            for col in range(1, len(headers) + 1):
                ws.cell(row=row, column=col).border = thin_border
                ws.cell(row=row, column=col).alignment = Alignment(horizontal="center")
            row += 1

    # Auto-width
    for col in ws.columns:
        max_len = max((len(str(cell.value or "")) for cell in col), default=10)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 20)

    # ── Per-pig sheets ──
    for pid in pig_ids:
        pig_results = sorted(
            [r for r in results if r["pig_id"] == pid],
            key=lambda r: r["pig_timestamp"],
        )
        if not pig_results:
            continue

        ws = wb.create_sheet(title=f"Pig {pid}")
        headers = ["Date", "Time", "Posture", "Confidence %", "Median Depth (mm)"]
        for col, h in enumerate(headers, 1):
            ws.cell(row=1, column=col, value=h)
        style_header(ws, 1, len(headers))

        for i, r in enumerate(pig_results, start=2):
            t = parse_timestamp(r["pig_timestamp"])
            ws.cell(row=i, column=1, value=t.strftime("%Y-%m-%d") if t else "")
            ws.cell(row=i, column=2, value=t.strftime("%H:%M:%S") if t else "")
            posture_cell = ws.cell(row=i, column=3, value=r["prediction_name"].capitalize())
            ws.cell(row=i, column=4, value=round(r["confidence"] * 100, 1))
            ws.cell(row=i, column=5, value=r["median_depth"])

            # Color-code posture cell
            fill = posture_fills.get(r["prediction_name"])
            if fill:
                posture_cell.fill = fill

            for col in range(1, len(headers) + 1):
                ws.cell(row=i, column=col).border = thin_border
                ws.cell(row=i, column=col).alignment = Alignment(horizontal="center")

        for col in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 20)

    wb.save(out_path)
    print(f"  Excel report: {out_path}")


def generate_timestamp_summary(results, out_path):
    """Generate per-timestamp-folder summary CSV."""
    from collections import defaultdict
    folder_stats = defaultdict(lambda: {"standing": 0, "sitting": 0, "lying": 0, "total": 0, "conf_sum": 0.0})

    for r in results:
        folder = r["timestamp_folder"]
        name = r["prediction_name"]
        folder_stats[folder][name] += 1
        folder_stats[folder]["total"] += 1
        folder_stats[folder]["conf_sum"] += r["confidence"]

    fields = ["timestamp_folder", "total_frames", "standing", "sitting", "lying",
              "standing_pct", "sitting_pct", "lying_pct", "avg_confidence"]

    with open(out_path, "w", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=fields)
        writer.writeheader()
        for folder in sorted(folder_stats):
            s = folder_stats[folder]
            t = s["total"]
            writer.writerow({
                "timestamp_folder": folder,
                "total_frames": t,
                "standing": s["standing"],
                "sitting": s["sitting"],
                "lying": s["lying"],
                "standing_pct": round(s["standing"] / t * 100, 1),
                "sitting_pct": round(s["sitting"] / t * 100, 1),
                "lying_pct": round(s["lying"] / t * 100, 1),
                "avg_confidence": round(s["conf_sum"] / t, 4),
            })

    print(f"  Timestamp summary: {out_path}")


def print_summary(results):
    counts = Counter(r["prediction_name"] for r in results)
    total = len(results)
    if total == 0:
        print("\n  No frames left after filtering (see predictions_excluded.csv / data_quality.csv)")
        return
    print(f"\n{'='*50}")
    print(f"  Total predicted: {total}")
    for name in ["standing", "sitting", "lying"]:
        c = counts.get(name, 0)
        print(f"  {name:<10}: {c:>6} ({c/total*100:.1f}%)")
    print(f"  Avg confidence: {np.mean([r['confidence'] for r in results]):.3f}")
    print(f"{'='*50}")


# ── Main ──────────────────────────────────────────────────────────────────

def load_posture_model(model_path, device):
    ckpt = torch.load(model_path, map_location=device, weights_only=False)
    backbone = ckpt.get("backbone", "mobilenet_v2")
    model = SingleModalModel(
        in_channels=ckpt.get("in_channels", 3),
        num_classes=ckpt.get("num_classes", 3),
        pretrained=False,
        backbone=backbone,
    ).to(device)
    model.load_state_dict(ckpt["model_state_dict"])
    model.eval()
    normalize, crop = ckpt_preprocessing(ckpt)
    model.posture_crop = crop
    return model, ckpt, normalize


def main():
    parser = argparse.ArgumentParser(
        description="One-command posture inference: input path → CSV + heatmap"
    )
    parser.add_argument("data_dir", help="Path to OAK-D data folder")
    parser.add_argument("--output-dir", default=None,
                        help="Output directory (default: outputs/)")
    parser.add_argument("--model", default=os.path.join(MODEL_DIR, "posture3_ir_best.pth"),
                        help="Posture checkpoint (default: models/posture3_ir_best.pth)")
    parser.add_argument("--batch-size", type=int, default=None,
                        help="Frames per forward pass (default: 64 on GPU/MPS, 4 on CPU — "
                             "large batches are slower on CPU for MobileNetV2)")
    parser.add_argument("--workers", type=int, default=None,
                        help="Image loader worker processes (default 0 on macOS/Windows, where spawning "
                             "workers made a 1k-frame run 3x slower; up to 4 on Linux)")
    parser.add_argument("--id-policy", choices=["strict", "modulo"], default="strict",
                        help="strict (default): apply known corrections, drop overflow IDs and "
                             "duplicate captures; modulo: v1.0 behaviour (pid %% 20)")
    parser.add_argument("--smooth", type=int, default=0,
                        help="Add smoothed_prediction_name using a centered majority vote over "
                             "N frames per pig (0/1 = off, 3 recommended)")
    parser.add_argument("--use-smoothed", action="store_true",
                        help="Use the smoothed posture for heatmaps/Excel instead of the raw one")
    parser.add_argument("--max-gap", type=float, default=30.0,
                        help="Max minutes between frames to treat them as consecutive")
    parser.add_argument("--no-detector", action="store_true",
                        help="Skip the CNN pig-presence detector (depth prefilter only)")
    args = parser.parse_args()

    data_dir = args.data_dir
    out_dir = args.output_dir or OUTPUT_DIR

    if not os.path.isdir(data_dir):
        print(f"Error: {data_dir} is not a directory")
        sys.exit(1)

    os.makedirs(out_dir, exist_ok=True)
    device = pick_device()
    if args.batch_size is None:
        args.batch_size = 4 if device.type == "cpu" else 64
    if args.workers is None:
        args.workers = max(0, min(4, (os.cpu_count() or 1) - 1)) if sys.platform.startswith("linux") else 0
    print(f"Device: {device}")

    if not os.path.exists(args.model):
        print(f"Error: model not found at {args.model}")
        sys.exit(1)
    model, ckpt, posture_norm = load_posture_model(args.model, device)
    print(f"Posture model: {os.path.basename(args.model)} "
          f"({ckpt.get('backbone', 'mobilenet_v2')}, {ckpt.get('num_classes', 3)} classes, "
          f"normalize={posture_norm})")

    detector = None if args.no_detector else load_pig_detector(MODEL_DIR, device)
    if detector is not None:
        print("Pig detector: MobileNetV2 (binary, present/empty)")
    else:
        print("Pig detector: not used — depth prefilter only")
    print()

    # Step 1: scan + ID policy on every capture
    frames = scan_folder(data_dir)
    if not frames:
        print("No frames found.")
        return
    csv_path = os.path.join(out_dir, "predictions.csv")
    excluded_path = os.path.join(out_dir, "predictions_excluded.csv")
    quality_path = os.path.join(out_dir, "data_quality.csv")
    coverage_path = os.path.join(out_dir, "coverage_by_pig_day.csv")
    matrix_path = os.path.join(out_dir, "posture_matrix.csv")
    heatmap_path = os.path.join(out_dir, "posture_heatmap.png")
    heatmap_raw_path = os.path.join(out_dir, "posture_heatmap_raw_ids.png")
    ts_summary_path = os.path.join(out_dir, "timestamp_summary.csv")

    # Quality report describes the raw captures (incl. double captures), written first
    quality = folder_quality(assign_ids([dict(f) for f in frames], policy=args.id_policy))
    save_rows(quality, quality_path, list(quality[0].keys()))
    # For the outputs, duplicates are resolved later among captures that produced a
    # prediction, so a corrupt/empty first capture does not hide a usable second one.
    assign_ids(frames, policy=args.id_policy, dedupe_captures=False, verbose=True)
    status_by_key = {(f["timestamp_folder"], f["pig_id_raw"], f["pig_timestamp"]): f for f in frames}

    # Step 2: predict (cropping happens in memory, no disk writes)
    raw_frames = [dict(f, pig_id=f["pig_id_raw"]) for f in frames]
    results, stats = run_predictions(data_dir, device, model, detector=detector,
                                     batch_size=args.batch_size, workers=args.workers,
                                     posture_normalize=posture_norm, frames=raw_frames,
                                     posture_crop=getattr(model, "posture_crop", True))
    if not results:
        print("No predictions generated.")
        return

    for r in results:
        f = status_by_key[(r["timestamp_folder"], r["pig_id"], r["pig_timestamp"])]
        r["pig_id_raw"] = f["pig_id_raw"]
        r["id_status"] = f["id_status"]
        r["include"] = f["include"]
        r["pig_id"] = f["pig_id"]
    dedupe(results)

    kept = [r for r in results if r["include"]]
    excluded = [r for r in results if not r["include"]]
    add_change_column(kept, max_gap_minutes=args.max_gap)
    if args.smooth > 1:
        smooth_results(kept, window=args.smooth, max_gap_minutes=args.max_gap)
        add_change_column(kept, key="smoothed_prediction_name",
                          out_key="posture_changed_smoothed", max_gap_minutes=args.max_gap)

    # Step 3: save outputs
    fields = PRED_FIELDS + (SMOOTH_FIELDS if args.smooth > 1 else [])
    save_rows(kept, csv_path, fields)
    save_rows(excluded, excluded_path, fields)
    save_coverage(kept, coverage_path)
    save_posture_matrix(kept, matrix_path,
                        key="smoothed_prediction_name" if (args.use_smoothed and args.smooth > 1) else "prediction_name")
    if not kept:
        print_summary(kept)
        print(f"  Every prediction was set aside by the ID policy — see {excluded_path} and {quality_path}")
        return
    print(f"\n[2/2] Generating outputs...")

    # Raw-ID heatmap shows every capture the camera produced (overflow visible)
    generate_heatmap_raw([dict(r, pig_id=r["pig_id_raw"]) for r in results], heatmap_raw_path)

    outputs = kept
    if args.use_smoothed and args.smooth > 1:
        outputs = [dict(r, prediction_name=r["smoothed_prediction_name"]) for r in kept]

    generate_heatmap(outputs, heatmap_path)
    generate_per_pig_heatmaps(outputs, out_dir)
    generate_per_pig_csvs(outputs, out_dir)
    generate_per_pig_excels(outputs, out_dir)
    excel_path = os.path.join(out_dir, "posture_report.xlsx")
    generate_excel_report(outputs, excel_path)
    generate_timestamp_summary(outputs, ts_summary_path)

    status_counts = Counter(r["id_status"] for r in results)
    n_review = sum(q["needs_review"] for q in quality)
    print_summary(outputs)
    print(f"  ID policy: {args.id_policy} | predicted-frame status: {dict(status_counts)}")
    print(f"  Folders needing review: {n_review}/{len(quality)} (see data_quality.csv)")
    print(f"\n  CSV:              {csv_path}")
    print(f"  Excluded rows:    {excluded_path}")
    print(f"  Data quality:     {quality_path}")
    print(f"  Coverage:         {coverage_path}")
    print(f"  Posture matrix:   {matrix_path}  (rows = folders, columns = pig0..pig19)")
    print(f"  Heatmap (all):    {heatmap_path}")
    print(f"  Heatmap (raw IDs):{heatmap_raw_path}")
    print(f"  Per-pig outputs:  {os.path.join(out_dir, 'per_pig')}/")
    print(f"  Excel report:     {excel_path}")
    print(f"  Timestamp summary:{ts_summary_path}")


if __name__ == "__main__":
    main()
