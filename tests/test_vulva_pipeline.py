"""Fast checks for the vulva pipeline (no lab data needed).

Run:  python tests/test_vulva_pipeline.py      (or: pytest tests/)

Synthetic OAK-D frames: a tilted rump plane in depth with an ellipsoid bump of
known width / length / height, and the matching IR image with the vulva drawn
in. Checks the 3D measurement against the known size, the readers for Lucas's
workbook and the manual labels, frame selection, daily aggregation, and the
whole train -> run path on a tiny dataset.
"""

import csv
import os
import shutil
import subprocess
import sys
import tempfile
from datetime import datetime

import cv2
import numpy as np

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "vulva"))
sys.path.insert(0, ROOT)

from daily import add_baseline_change, aggregate, compare_truth  # noqa: E402
from measure3d import DEFAULT_INTRINSICS as K  # noqa: E402
from measure3d import measure  # noqa: E402
from select_frames import parse_windows, select  # noqa: E402

PY = sys.executable


def render(tilt_deg=30, dist=1200, W=40.0, L=60.0, H=15.0, noise=0.0, seed=0, cx=320, cy=260):
    """Depth (uint16 mm) and vulva mask for a rump plane tilted about x with an ellipsoid bump."""
    t = np.radians(tilt_deg)
    n = np.array([0, np.sin(t), -np.cos(t)])
    ex = np.array([1.0, 0, 0])
    ey = np.cross(n, ex)
    ey /= np.linalg.norm(ey)
    if ey[1] < 0:
        ey = -ey
    r0 = np.array([(cx - K["cx"]) / K["fx"], (cy - K["cy"]) / K["fy"], 1.0])
    C = r0 / r0[2] * dist
    ys, xs = np.mgrid[0:480, 0:640]
    rays = np.stack([(xs - K["cx"]) / K["fx"], (ys - K["cy"]) / K["fy"], np.ones_like(xs, float)], -1)
    s = (C @ n) / (rays @ n)
    P = rays * s[..., None]
    rel = P - C
    a, b = rel @ ex, rel @ ey
    q = 1 - (a / (W / 2)) ** 2 - (b / (L / 2)) ** 2
    h = np.where(q > 0, H * np.sqrt(np.clip(q, 0, 1)), 0.0)
    unit = rays / np.linalg.norm(rays, axis=-1, keepdims=True)
    cos = np.abs(unit @ n)
    z = P[..., 2] - h / cos * unit[..., 2]
    z = z + np.random.default_rng(seed).normal(0, noise, z.shape)
    return np.clip(z, 0, 65535).astype(np.uint16), q > 0


def test_measure_recovers_known_size_under_tilt():
    for tilt in (0, 45, 60):
        depth, mask = render(tilt_deg=tilt)
        m = measure(depth, mask)
        assert m["measure_status"] == "ok"
        assert abs(m["width_mm"] - 40) < 2.0, (tilt, m["width_mm"])
        assert abs(m["length_mm"] - 60) < 2.5, (tilt, m["length_mm"])
        assert abs(m["peak_height_mm"] - 15) < 1.5, (tilt, m["peak_height_mm"])
        true_vol = 2 / 3 * np.pi * 20 * 30 * 15
        assert abs(m["volume_mm3"] - true_vol) / true_vol < 0.10, (tilt, m["volume_mm3"])
        assert abs(m["view_angle_deg"] - tilt) < 1.0


def test_measure_handles_noise_bad_masks_and_missing_depth():
    depth, mask = render(tilt_deg=30, noise=2.0, seed=3)
    m = measure(depth, mask)
    assert abs(m["width_mm"] - 40) < 3 and abs(m["length_mm"] - 60) < 4
    assert measure(depth, np.zeros_like(mask))["measure_status"] == "empty_mask"
    assert measure(None, mask)["measure_status"] == "no_depth"
    tiny = np.zeros_like(mask)
    tiny[260:263, 320:323] = True
    assert measure(depth, tiny)["measure_status"] == "too_few_points"


def _lucas_workbook(path):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    days = [(0, "MOVED", None, None, 0, 0), (1, "OK", 40.0, 60.0, 0, 0), (2, "", 40.4, 60.6, 0, 0),
            (3, "OK", 44.0, 66.0, 20, 20), (4, "", 43.0, 65.0, 0, 20), (5, "OK", 40.2, 60.1, 0, 0),
            (6, "", 40.0, 60.0, 0, 0), (7, "OK", 41.0, 59.0, 0, 0)]
    for pid, tag in ((0, "200196"), (1, "182")):
        ws = wb.create_sheet(str(pid))
        ws["A1"], ws["A2"], ws["B2"] = "SOW 20", "NUMBER", int(tag)
        ws["A13"], ws["B13"], ws["C13"], ws["H13"], ws["P13"] = "DAY", "DATE", "BODY CONDITION", "VULVA", "HEAT"
        for col, name in zip("CDEFGHIJKLMNOPQ", ["CALIPER", "BS", "BC SOLID", "WEIGHT", "BACKFAT", "V. W.", "V. L.",
                                                  "V. V.", "V. S. O. A.", "V. S. N. A.", "V. V. S. B. N. E.",
                                                  "I. I. T. B", "BPTM", "MORNING", "EVENING"]):
            ws[f"{col}14"] = name
        ws["S12"] = "ACTION"
        for i, (day, action, w, l, hm, he) in enumerate(days):
            r = 15 + i
            ws[f"A{r}"], ws[f"B{r}"] = day, datetime(2026, 2, 6 + day)
            if w:
                ws[f"C{r}"], ws[f"D{r}"], ws[f"H{r}"], ws[f"I{r}"] = 14.5, 4, w, l
                ws[f"J{r}"] = w * l / 200
            ws[f"P{r}"], ws[f"Q{r}"], ws[f"S{r}"] = hm, he, action
        ws["A40"], ws["B40"], ws["J40"] = 30, datetime(2026, 3, 8), 0     # empty future row
    wb.save(path)


def test_truth_reads_lucas_workbook_layout():
    from truth import heat_response, read_workbook
    tmp = tempfile.mkdtemp()
    try:
        p = os.path.join(tmp, "Feb_2026.xlsx")
        _lucas_workbook(p)
        rows = read_workbook(p)
        assert len(rows) == 16 and {r["pig_id"] for r in rows} == {0, 1}
        r3 = next(r for r in rows if r["pig_id"] == 0 and r["day"] == 3)
        assert r3["date"] == "2026-02-09" and r3["heat"] == 1 and r3["vulva_width_mm"] == 44.0
        assert abs(r3["vulva_index"] - 44 * 66 / 200) < 1e-6
        assert next(r for r in rows if r["day"] == 0)["action"] == "MOVED"
        assert rows[0]["tag"] == "200196"
        summary, noise = heat_response(rows)
        s0 = summary[0]
        assert s0["heat_days"] == "3|4"
        assert abs(s0["vulva_index_rise_pct"] - 100 * (44 * 66 - 40.2 * 60.3) / (40.2 * 60.3)) < 1.5
        assert noise["vulva_width_mm"] is not None
    finally:
        shutil.rmtree(tmp)


def test_select_frames_standing_windows_and_cap():
    rows = []
    for h in range(0, 24):
        for pid in (0, 1):
            rows.append({"timestamp_folder": f"20260211-{h:02d}-00-00", "pig_id": str(pid),
                         "pig_timestamp": f"20260211-{h:02d}-00-{10 + pid:02d}",
                         "prediction_name": "standing" if h % 2 == 0 else "lying",
                         "confidence": "0.9", "id_status": "ok",
                         "image_path": f"/d/20260211-{h:02d}-00-00/pig{pid}_ir_vis_20260211-{h:02d}-00-{10 + pid:02d}.jpg"})
    rows.append(dict(rows[0], id_status="overflow"))
    picked = select(rows, windows=parse_windows("06:00-11:00"))
    assert {p["pig_timestamp"][9:11] for p in picked} == {"06", "08", "10"}
    assert picked[0]["depth_path"].endswith("pig0_depth_20260211-06-00-10.raw")
    capped = select(rows, max_per_day=4)
    assert len(capped) == 8 and len({p["pig_timestamp"] for p in capped if p["pig_id"] == 0}) == 4
    one = select(rows, max_per_day=1)                   # 12 standing hours: the middle one, not the first
    assert [p["pig_timestamp"][9:11] for p in one if p["pig_id"] == 0] == ["12"]


def test_relocate_finds_frames_after_the_data_moved():
    """Predictions made before the raw folders moved: frames are found again by timestamp folder name."""
    from select_frames import relocate
    tmp = tempfile.mkdtemp()
    try:
        new = os.path.join(tmp, "E", "Pictures_OAK", "OAK_0206_0211", "20260209-09-00-00")
        os.makedirs(new)
        open(os.path.join(new, "pig3_ir_vis_20260209-09-00-13.jpg"), "wb").close()
        picked = [{"timestamp_folder": "20260209-09-00-00", "pig_id": 3, "pig_timestamp": "20260209-09-00-13",
                   "ir_path": "C:\\PAAL_data\\Fed_pig\\Pictures_OAK\\20260209-09-00-00\\pig3_ir_vis_20260209-09-00-13.jpg",
                   "depth_path": ""},
                  dict(timestamp_folder="20260209-09-10-00", pig_id=4, pig_timestamp="20260209-09-10-14",
                       ir_path="C:\\gone\\pig4_ir_vis_20260209-09-10-14.jpg", depth_path="")]
        found, missing = relocate(picked, [os.path.join(tmp, "E", "Pictures_OAK")])
        assert missing == 1 and len(found) == 1
        assert found[0]["ir_path"] == os.path.join(new, "pig3_ir_vis_20260209-09-00-13.jpg")
        assert found[0]["depth_path"].endswith("pig3_depth_20260209-09-00-13.raw")
    finally:
        shutil.rmtree(tmp)


def test_review_run_panels_and_point_cloud():
    """review_run.py: panels for a run's frames, and the 3D surface / point cloud around a vulva."""
    from review_run import local_surface, write_ply
    from seg_model import VulvaUNet, save_checkpoint
    depth, mask = render(tilt_deg=30, W=40, L=60, H=12, noise=0.0)
    surf = local_surface(depth, mask, 320, 260)
    assert surf is not None
    u, v, h, vul, pts, ys, xs = surf
    assert 8 < np.percentile(h[vul], 95) < 16 and abs(np.median(h[~vul])) < 2      # bump on a flat rump
    tmp = tempfile.mkdtemp()
    try:
        write_ply(os.path.join(tmp, "a.ply"), surf, np.full((480, 640), 128, np.uint8))
        with open(os.path.join(tmp, "a.ply")) as f:
            head = f.read(400)
        assert head.startswith("ply") and f"element vertex {len(pts)}" in head
        rng = np.random.default_rng(2)
        folder = os.path.join(tmp, "20260124-09-00-00")
        os.makedirs(folder)
        rows = []
        for pid in range(2):
            ts = f"20260124-09-00-{10 + pid:02d}"
            _write_frame(folder, pid, ts, True, rng)
            rows.append({"timestamp_folder": "20260124-09-00-00", "pig_id": pid, "pig_timestamp": ts,
                         "measure_status": "ok", "width_mm": 40, "length_mm": 60,
                         "ir_path": os.path.join(folder, f"pig{pid}_ir_vis_{ts}.jpg")})
        fcsv = os.path.join(tmp, "vulva_frames.csv")
        with open(fcsv, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        ck = os.path.join(tmp, "m.pth")
        save_checkpoint(VulvaUNet(pretrained=False), ck, {"threshold": 0.5, "min_area": 25, "inputs": "ir"})
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "review_run.py"), fcsv, "--model", ck, "--n", "2"],
                             capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        assert len(os.listdir(os.path.join(tmp, "review_ok", "panels"))) == 2
        assert cv2.imread(os.path.join(tmp, "review_ok", "index.jpg")).shape[1] == 600
    finally:
        shutil.rmtree(tmp)


def test_daily_median_baseline_and_caliper_join():
    frames = []
    for day, base in ((7, 40.0), (8, 40.0), (9, 44.0), (10, 40.0)):
        for k in range(5):
            frames.append({"timestamp_folder": f"202602{day:02d}-09-{k:02d}-00", "pig_id": "3",
                           "measure_status": "ok", "width_mm": base + (k - 2) * 0.5, "length_mm": 60.0,
                           "peak_height_mm": 10, "volume_mm3": 1000, "plane_rmse_mm": 1, "view_angle_deg": 30,
                           "depth_coverage": 0.9})
    frames.append(dict(frames[0], width_mm=99, plane_rmse_mm=50))     # bad plane fit, dropped
    daily = add_baseline_change(aggregate(frames, min_frames=3), baseline_days=2)
    assert [d["n_frames"] for d in daily] == [5, 5, 5, 5]
    assert daily[0]["width_mm"] == 40.0 and daily[2]["width_mm_change_pct"] == 10.0
    truth = [{"pig_id": 3, "date": f"2026-02-{d:02d}", "vulva_width_mm": w, "vulva_length_mm": 60,
              "vulva_index": w * 60 / 200, "heat": int(d == 9)} for d, w in ((7, 41), (8, 41), (9, 45), (10, 41))]
    joined, s = compare_truth(daily, truth)
    assert s["sow_days"] == 4 and s["width_mm_mae"] == 1.0 and s["width_mm_bias"] == -1.0
    assert s["index_change_heat_days"] > s["index_change_other_days"]


def test_repeatability_reports_frames_needed():
    from repeatability import analyse
    rows = [{"timestamp_folder": "20260211-09-00-00", "pig_id": "1", "status": "measured",
             "width_mm": w, "length_mm": 60} for w in (36, 40, 44, 40)]
    per, s = analyse(rows, truth=[{"pig_id": 1, "date": "2026-02-11", "vulva_width_mm": 40, "vulva_length_mm": 60}])
    assert s["width_mm_within_day_sd"] > 3 and s["width_mm_frames_needed"] == 3
    assert s["width_mm_mae_vs_caliper"] == 2.0 and s["length_mm_mae_vs_caliper"] == 0.0


def _write_frame(folder, pid, ts, present, rng):
    """IR jpg + depth raw for one stall; returns the full-frame mask when present."""
    tilt = rng.uniform(20, 50)
    W, L = rng.uniform(34, 46), rng.uniform(52, 70)
    cx, cy = int(rng.uniform(280, 360)), int(rng.uniform(220, 300))
    depth, mask = render(tilt_deg=tilt, W=W, L=L, H=14, cx=cx, cy=cy, noise=1.0, seed=int(rng.integers(1e6)))
    ir = np.clip(rng.normal(110, 12, (480, 640)), 0, 255).astype(np.uint8)
    ir = cv2.GaussianBlur(ir, (5, 5), 0)
    if present:
        ir[mask] = 35
        ring = cv2.dilate(mask.astype(np.uint8), np.ones((5, 5), np.uint8)) > 0
        ir[ring & ~mask] = 200
    else:
        mask = np.zeros_like(mask)
    cv2.imwrite(os.path.join(folder, f"pig{pid}_ir_vis_{ts}.jpg"), ir)
    with open(os.path.join(folder, f"pig{pid}_depth_{ts}.raw"), "wb") as f:
        f.write(b"\0" * 8)
        f.write(depth.tobytes())
    return mask


def test_train_and_run_end_to_end():
    tmp = tempfile.mkdtemp()
    try:
        rng = np.random.default_rng(0)
        data = os.path.join(tmp, "Pictures_OAK")
        masks = os.path.join(tmp, "labels", "vulva_masks_manual")
        os.makedirs(masks)
        manual, preds = [], []
        for k in range(6):
            folder_name = f"20260211-09-{10 * k:02d}-00"
            folder = os.path.join(data, folder_name)
            os.makedirs(folder)
            for pid in range(4):
                ts = f"20260211-09-{10 * k:02d}-{10 + pid:02d}"
                present = (k + pid) % 3 != 0
                m = _write_frame(folder, pid, ts, present, rng)
                mp = ""
                if present:
                    mp = os.path.join("C:\\elsewhere\\labels\\vulva_masks_manual",
                                      f"{folder_name}_pig{pid}_{ts}_vulva_mask.png")
                    cv2.imwrite(os.path.join(masks, os.path.basename(mp.replace("\\", "/"))), m.astype(np.uint8) * 255)
                manual.append({"timestamp_folder": folder_name, "pig_id": pid, "pig_timestamp": ts,
                               "status": "measured" if present else "no_vulva", "mask_path": mp,
                               "texture_input": ""})
                preds.append({"timestamp_folder": folder_name, "pig_id": pid, "pig_timestamp": ts,
                              "prediction_name": "standing", "confidence": 0.95, "id_status": "ok",
                              "image_path": os.path.join(folder, f"pig{pid}_ir_vis_{ts}.jpg")})
        manual.append(dict(manual[0], status="failed: Vulva mask is empty after cleanup"))
        man_csv = os.path.join(tmp, "manual.csv")
        with open(man_csv, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(manual[0]))
            w.writeheader()
            w.writerows(manual)
        pred_csv = os.path.join(tmp, "predictions.csv")
        with open(pred_csv, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(preds[0]))
            w.writeheader()
            w.writerows(preds)

        labels = os.path.join(tmp, "seg_labels.csv")
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "make_seg_labels.py"), "--manual", man_csv,
                              "--data-root", data, "--mask-root", masks, "--out", labels],
                             capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr
        assert "24 labelled frames (16 present, 8 absent)" in out.stdout, out.stdout

        model = os.path.join(tmp, "models", "vulva_seg.pth")
        share = os.path.join(tmp, "share")
        os.makedirs(share)
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "train_seg.py"), "--labels", labels, "--folds", "2",
                              "--epochs", "4", "--no-pretrained", "--out-model", model,
                              "--out-dir", os.path.join(tmp, "seg"), "--copy-to", share],
                             capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        assert os.path.exists(model) and "Cross-validation" in out.stdout
        assert os.path.exists(os.path.join(tmp, "seg", "seg_report.json"))
        assert sorted(os.listdir(share)) == ["vulva_seg.pth", "vulva_seg_cv.csv", "vulva_seg_report.txt"]
        # --jobs: folds and final model as parallel workers, same outputs
        model_p = os.path.join(tmp, "models", "vulva_seg_par.pth")
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "train_seg.py"), "--labels", labels, "--folds", "2",
                              "--epochs", "3", "--no-pretrained", "--jobs", "3", "--out-model", model_p,
                              "--out-dir", os.path.join(tmp, "seg_par")], capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stdout[-2000:] + out.stderr[-2000:]
        assert "All 3 models trained" in out.stdout and "Cross-validation" in out.stdout, out.stdout
        assert os.path.exists(model_p)
        with open(os.path.join(tmp, "seg_par", "cv_predictions.csv")) as f:
            assert len(list(csv.DictReader(f))) == 24
        out_train = out.stdout

        outdir = os.path.join(tmp, "out")
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "run_vulva.py"), pred_csv, "--model", model,
                              "--out-dir", outdir, "--min-frames", "1", "--save-masks"],
                             capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        with open(os.path.join(outdir, "vulva_frames.csv")) as f:
            rows = list(csv.DictReader(f))
        assert len(rows) == 24
        assert all(r["measure_status"] for r in rows)
        assert "vulva_conf" in rows[0]
        with open(os.path.join(tmp, "seg", "cv_predictions.csv")) as f:
            assert f.readline().strip().endswith("peak_prob,blob_conf")
        assert "With confidence cut" in out_train or "No confidence cut" in out_train
        assert os.path.exists(os.path.join(outdir, "vulva_daily.csv"))
        assert os.path.exists(os.path.join(outdir, "vulva_summary.json"))
        # IR + depth model: trains, records its inputs, and run_vulva feeds it depth
        model_d = os.path.join(tmp, "models", "vulva_seg_d.pth")
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "train_seg.py"), "--labels", labels, "--folds", "0",
                              "--epochs", "3", "--no-pretrained", "--inputs", "ir+depth", "--out-model", model_d,
                              "--out-dir", os.path.join(tmp, "seg_d")], capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        assert "(inputs: ir+depth)" in out.stdout
        import torch
        assert torch.load(model_d, map_location="cpu", weights_only=False)["inputs"] == "ir+depth"
        assert "presence_threshold" in torch.load(model, map_location="cpu", weights_only=False)
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "run_vulva.py"), pred_csv, "--model", model_d,
                              "--out-dir", os.path.join(tmp, "out_d"), "--min-frames", "1"],
                             capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        # nothing written next to the raw frames
        assert all(not n.endswith(".png") for _, _, fs in os.walk(data) for n in fs)
    finally:
        shutil.rmtree(tmp)


def test_ellipse_from_three_clicks():
    from label_ellipse import ellipse_from_clicks, ellipse_mask
    e = ellipse_from_clicks((300, 200), (300, 260), (320, 230))      # vertical, 60 long, 40 wide
    assert e[0] == 300 and e[1] == 230 and e[2] == 30 and abs(e[3] - 20) < 1e-9 and abs(e[4] - 90) < 1e-9
    m = ellipse_mask(e) > 0
    ys, xs = np.nonzero(m)
    assert ys.min() == 200 and ys.max() == 260 and xs.min() == 280 and xs.max() == 320
    assert abs(m.sum() - np.pi * 30 * 20) / (np.pi * 30 * 20) < 0.05


def test_model_mask_to_clicks_round_trip():
    """A predicted mask becomes three clicks that redraw the same ellipse (model-assisted labelling)."""
    from label_ellipse import clicks_from_mask, ellipse_from_clicks, ellipse_mask
    assert clicks_from_mask(np.zeros((480, 640), np.uint8)) is None
    for ang in (90, 75, 110, 30, 0):
        truth = (330.0, 240.0, 40.0, 22.0, float(ang))
        clicks = clicks_from_mask(ellipse_mask(truth))
        top, bottom, _ = clicks
        assert top[1] <= bottom[1]
        e = ellipse_from_clicks(*clicks)
        assert abs(e[0] - 330) < 1.5 and abs(e[1] - 240) < 1.5, (ang, e)
        assert abs(e[2] - 40) < 1.5 and abs(e[3] - 22) < 1.5, (ang, e)
        a = ellipse_mask(truth) > 0
        b = ellipse_mask(e) > 0
        assert (a & b).sum() / (a | b).sum() > 0.93, (ang, e)


def test_presence_cut_reaches_target_precision():
    """The 'vulva visible' cut is the lowest blob confidence whose precision reaches the target."""
    from train_seg import pick_presence_cut
    vis, none = np.ones((4, 4), np.uint8), np.zeros((4, 4), np.uint8)
    items = [{"mask": vis}] * 7 + [{"mask": none}] * 4             # 7 real vulvas, 4 frames without
    pred = [vis] * 6 + [none] + [vis] * 4                           # one missed, four false blobs
    confs = [.9, .8, .7, .6, .5, .4, 0, .65, .45, .35, .3]
    assert pick_presence_cut(items, pred, confs, 0.8) == 0.5        # 5 real, 1 false kept
    assert pick_presence_cut(items, pred, confs, 1.0) == 0.7
    assert pick_presence_cut(items, pred, confs, 0) == 0.0
    assert pick_presence_cut(items, pred, [.1] * 6 + [0] + [.95] * 4, 0.9) is None


def test_saved_ellipse_round_trip_for_review():
    """Review mode redraws a saved ellipse from its stored parameters; Enter must save the same one."""
    from label_ellipse import ellipse_from_clicks, saved_clicks
    for ang in (90.0, 75.0, 110.0, -80.0):
        row = {"cx": "330", "cy": "250", "half_length_px": "38", "half_width_px": "21", "angle_deg": str(ang)}
        e = ellipse_from_clicks(*saved_clicks(row))
        assert abs(e[0] - 330) < 1e-6 and abs(e[1] - 250) < 1e-6 and abs(e[2] - 38) < 1e-6 and abs(e[3] - 21) < 1e-6
        assert abs(((e[4] - ang) + 90) % 180 - 90) < 1e-6          # same axis (direction may flip)
    assert saved_clicks({"cx": "", "cy": ""}) is None


def test_review_mode_keeps_or_changes_saved_labels():
    """--review: Enter keeps the saved ellipse (how unchanged), t turns it into 'tail' (not visible)."""
    tmp = tempfile.mkdtemp()
    try:
        rng = np.random.default_rng(3)
        folder = os.path.join(tmp, "20260209-09-00-00")
        os.makedirs(folder)
        fields = ["timestamp_folder", "pig_id", "pig_timestamp", "ir_path", "depth_path", "posture_confidence"]
        frames, labels = [], []
        for pid in range(3):
            ts = f"20260209-09-00-{10 + pid:02d}"
            _write_frame(folder, pid, ts, True, rng)
            ir = f"20260209-09-00-00/pig{pid}_ir_vis_{ts}.jpg"
            frames.append({"timestamp_folder": "20260209-09-00-00", "pig_id": pid, "pig_timestamp": ts, "ir_path": ir,
                           "depth_path": "", "posture_confidence": 0.9})
            labels.append({"ir_path": ir, "label": "present", "pig_id": pid,
                           "timestamp_folder": "20260209-09-00-00", "pig_timestamp": ts, "cx": 320, "cy": 250,
                           "half_length_px": 30, "half_width_px": 18, "angle_deg": 90, "how": "accepted",
                           "mask_path": f"masks/old{pid}.png"})
        for name, rows in (("label_set.csv", frames), ("labels.csv", labels)):
            with open(os.path.join(tmp, name), "w", newline="") as f:
                w = csv.DictWriter(f, fieldnames=list(rows[0]))
                w.writeheader()
                w.writerows(rows)
        driver = os.path.join(tmp, "drive.py")
        with open(driver, "w") as f:
            f.write(
                "import sys, matplotlib\nmatplotlib.use('Agg')\nimport matplotlib.pyplot as plt\n"
                "from matplotlib.backend_bases import KeyEvent\n"
                "def show():\n    fig = plt.gcf()\n"
                "    for k in ('enter', 't', 'q'):\n"
                "        fig.canvas.callbacks.process('key_press_event', KeyEvent('key_press_event', fig.canvas, k))\n"
                "plt.show = show\n"
                f"sys.path.insert(0, {os.path.join(ROOT, 'vulva')!r})\n"
                f"sys.argv = ['label_ellipse.py', {tmp!r}, '--review', 'r6', '--no-depth', '--no-rgb']\n"
                "import label_ellipse\nlabel_ellipse.main()\n")
        out = subprocess.run([PY, driver], capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        assert "Review r6: 3 to go" in out.stdout and "1 kept, 1 tail" in out.stdout, out.stdout
        with open(os.path.join(tmp, "labels.csv")) as f:
            rows = {r["pig_id"]: r for r in csv.DictReader(f)}
        assert rows["0"]["review"] == "r6 kept" and rows["0"]["how"] == "accepted" and rows["0"]["label"] == "present"
        assert rows["0"]["mask_path"] == "masks/old0.png"                  # kept: the saved mask is not rewritten
        assert rows["1"]["label"] == "absent" and rows["1"]["note"] == "tail" and rows["1"]["review"] == "r6 tail"
        assert rows["2"]["review"] == "" and rows["2"]["label"] == "present"
    finally:
        shutil.rmtree(tmp)


def test_rgb_align_recovers_a_known_shift():
    """rgb_align: a colour image made from the IR with a known offset is put back on the IR grid;
    register_rgb.py writes it into a label set without touching frames or labels."""
    from rgb_align import MAP, register
    from register_rgb import register_set
    rng = np.random.default_rng(11)
    ir = cv2.GaussianBlur(rng.integers(0, 255, (480, 640)).astype(np.uint8), (0, 0), 6)
    ir = cv2.normalize(ir, None, 0, 255, cv2.NORM_MINMAX)
    for _ in range(25):
        cv2.ellipse(ir, (int(rng.integers(120, 520)), int(rng.integers(80, 420))),
                    (int(rng.integers(8, 30)), int(rng.integers(8, 30))), 0, 0, 360, int(rng.integers(0, 255)), -1)
    s, dx, dy = MAP
    true = (21, 4)
    A = np.array([[1 / s, 0, -dx / s - true[0]], [0, 1 / s, -dy / s - true[1]]], np.float32)
    rgb = cv2.cvtColor(cv2.warpAffine(ir, A, (1280, 800), flags=cv2.INTER_LINEAR | cv2.WARP_INVERSE_MAP),
                       cv2.COLOR_GRAY2BGR)
    img, shift, score, how = register(ir, rgb, (330, 250))
    assert how == "vulva" and abs(shift[0] - true[0]) <= 1 and abs(shift[1] - true[1]) <= 1, (shift, score)
    back = cv2.cvtColor(img, cv2.COLOR_BGR2GRAY)[100:400, 150:500].astype(float)
    assert np.abs(back - ir[100:400, 150:500]).mean() < 6
    tmp = tempfile.mkdtemp()
    try:
        os.makedirs(os.path.join(tmp, "f"))
        cv2.imwrite(os.path.join(tmp, "f", "pig1_ir_vis_20260213-09-00-11.jpg"), ir)
        cv2.imwrite(os.path.join(tmp, "f", "pig1_rgb_20260213-09-00-11.jpg"), rgb)
        with open(os.path.join(tmp, "label_set.csv"), "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["timestamp_folder", "pig_id", "pig_timestamp", "ir_path", "rgb_path"])
            w.writeheader()
            w.writerow({"timestamp_folder": "f", "pig_id": 1, "pig_timestamp": "20260213-09-00-11",
                        "ir_path": "f/pig1_ir_vis_20260213-09-00-11.jpg", "rgb_path": "f/pig1_rgb_20260213-09-00-11.jpg"})
        counts = register_set(tmp)
        assert counts["none"] == 0 and sum(counts.values()) == 1
        with open(os.path.join(tmp, "label_set.csv")) as f:
            row = next(csv.DictReader(f))
        assert row["rgb_aligned_path"] == "f/pig1_rgbreg_20260213-09-00-11.jpg"
        assert cv2.imread(os.path.join(tmp, row["rgb_aligned_path"])).shape == (480, 640, 3)
        assert os.path.exists(os.path.join(tmp, "label_set_before_rgbreg.csv"))
    finally:
        shutil.rmtree(tmp)


def test_depth_input_stack():
    """ir+depth input: channel 0 is the IR, channel 1 height (128 = level skin, higher on the bump),
    channel 2 shading; missing depth gives a flat, cue-free stack."""
    from seg_model import make_input, to_model_space, to_tensor
    depth, mask = render(tilt_deg=20, dist=700, W=40, L=70, H=15, noise=0.5, cx=320, cy=240)
    ir = np.full((480, 640), 90, np.uint8)
    x = make_input(ir, depth, "ir+depth")
    assert x.shape == (480, 640, 3) and (x[..., 0] == 90).all()
    assert np.median(x[..., 1][mask]) > 150 and abs(int(np.median(x[..., 1][380:450, 400:560])) - 128) < 10
    flat = make_input(ir, None, "ir+depth")
    assert (flat[..., 1] == 128).all() and (flat[..., 2] == 0).all()
    assert make_input(ir, depth, "ir") is ir
    im, _ = to_model_space(x)
    assert to_tensor(im).shape == (3, 384, 320) and to_tensor(to_model_space(ir)[0]).shape == (3, 384, 320)


def test_depth_views_show_the_bump_and_drop_the_background():
    """Height view: the vulva bump reads as several mm proud of the skin, flat skin near 0,
    and stall background far behind the sow is masked out (NaN)."""
    from depth_views import relief, shade_image, height_image
    depth, mask = render(tilt_deg=20, dist=700, W=40, L=70, H=15, noise=1.0, cx=320, cy=240)
    depth = depth.copy()
    depth[:, :150] = 2000                                    # far stall wall at the left
    shade, height, zc = relief(depth)
    assert 600 < zc < 850, zc
    assert np.isnan(height[:, :140]).all()                   # background is not part of the sow
    ys, xs = np.nonzero(mask)
    peak = np.nanmedian(height[ys.min() + (ys.max() - ys.min()) // 3: ys.max() - (ys.max() - ys.min()) // 3,
                               xs.min() + 5: xs.max() - 5])
    skin = np.nanmedian(np.abs(height[380:450, 400:560]))
    assert peak > 5 and skin < 1.5, (peak, skin)
    assert shade_image(shade).shape == (480, 640, 3) and height_image(height).dtype == np.uint8


def test_add_rgb_keeps_frames_and_labels():
    """--add-rgb copies pigN_rgb_aligned / pigN_rgb next to the frames already in the set and
    adds their columns; the frame list and labels.csv stay as they were."""
    from export_label_set import export, add_rgb
    tmp = tempfile.mkdtemp()
    try:
        rng = np.random.default_rng(9)
        raw = os.path.join(tmp, "raw", "20260210-09-00-00")
        os.makedirs(raw)
        preds = []
        for pid in range(3):
            ts = f"20260210-09-00-{10 + pid:02d}"
            _write_frame(raw, pid, ts, True, rng)
            cv2.imwrite(os.path.join(raw, f"pig{pid}_rgb_{ts}.jpg"), np.full((720, 1280, 3), 90, np.uint8))
            if pid != 1:                                   # one frame without an aligned image
                cv2.imwrite(os.path.join(raw, f"pig{pid}_rgb_aligned_{ts}.jpg"), np.full((480, 640, 3), 120, np.uint8))
            preds.append({"timestamp_folder": "20260210-09-00-00", "pig_id": pid, "pig_timestamp": ts,
                          "prediction_name": "standing", "confidence": 0.9, "id_status": "ok",
                          "image_path": os.path.join(raw, f"pig{pid}_ir_vis_{ts}.jpg")})
        out = os.path.join(tmp, "set")
        rows, _ = export(preds, out, per_day=3)
        assert "rgb_path" not in rows[0]
        with open(os.path.join(out, "labels.csv"), "w") as f:
            f.write("ir_path,label\n" + rows[0]["ir_path"] + ",absent\n")
        before = open(os.path.join(out, "label_set.csv")).read().splitlines()
        n, counts, no_source = add_rgb(preds, out)
        assert n == 3 and no_source == 0 and counts == {"rgb_aligned_path": 2, "rgb_path": 3}, counts
        with open(os.path.join(out, "label_set.csv"), newline="") as f:
            new = list(csv.DictReader(f))
        assert [r["ir_path"] for r in new] == [ln.split(",")[3] for ln in before[1:]]
        assert all(os.path.exists(os.path.join(out, r["rgb_path"])) for r in new)
        assert new[1]["rgb_aligned_path"] == "" and os.path.exists(os.path.join(out, new[0]["rgb_aligned_path"]))
        assert open(os.path.join(out, "labels.csv")).read().count("absent") == 1
        assert len(os.listdir(raw)) == 3 * 3 + 2                    # raw folder only read: ir, depth, rgb (+2 aligned)
    finally:
        shutil.rmtree(tmp)


def test_export_label_set_and_train_on_relative_paths():
    from export_label_set import export
    from train_seg import load_labels
    tmp = tempfile.mkdtemp()
    try:
        rng = np.random.default_rng(5)
        raw = os.path.join(tmp, "raw", "20260209-09-00-00")
        os.makedirs(raw)
        preds = []
        for pid in range(3):
            ts = f"20260209-09-00-{10 + pid:02d}"
            _write_frame(raw, pid, ts, True, rng)
            preds.append({"timestamp_folder": "20260209-09-00-00", "pig_id": pid, "pig_timestamp": ts,
                          "prediction_name": "standing" if pid < 2 else "lying", "confidence": 0.9,
                          "id_status": "ok", "image_path": os.path.join(raw, f"pig{pid}_ir_vis_{ts}.jpg")})
        out = os.path.join(tmp, "set")
        rows, missing = export(preds, out, per_day=3)
        assert len(rows) == 2 and missing == 0
        assert os.path.exists(os.path.join(out, rows[0]["depth_path"]))
        assert not os.path.isabs(rows[0]["ir_path"])
        # labels written the way label_ellipse.py writes them: relative to the set folder
        os.makedirs(os.path.join(out, "masks"))
        m = np.zeros((480, 640), np.uint8)
        m[250:270, 310:325] = 255
        cv2.imwrite(os.path.join(out, "masks", "a.png"), m)
        with open(os.path.join(out, "labels.csv"), "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=["ir_path", "mask_path", "label", "pig_id",
                                              "timestamp_folder", "pig_timestamp"])
            w.writeheader()
            w.writerow({"ir_path": rows[0]["ir_path"], "mask_path": "masks/a.png", "label": "present",
                        "pig_id": 0, "timestamp_folder": rows[0]["timestamp_folder"],
                        "pig_timestamp": rows[0]["pig_timestamp"]})
            w.writerow({"ir_path": rows[1]["ir_path"], "mask_path": "", "label": "absent", "pig_id": 1,
                        "timestamp_folder": rows[1]["timestamp_folder"], "pig_timestamp": rows[1]["pig_timestamp"]})
        items = load_labels(os.path.join(out, "labels.csv"))
        assert len(items) == 2 and items[0]["mask"].any() and not items[1]["mask"].any()
    finally:
        shutil.rmtree(tmp)


def test_eval_seg_scores_a_model_end_to_end():
    """eval_seg.py: a model against hand labels it was not trained on, with sizes in mm from both masks."""
    from eval_seg import evaluate, read_rows, summarise
    from seg_model import VulvaUNet
    tmp = tempfile.mkdtemp()
    try:
        rng = np.random.default_rng(7)
        folder = os.path.join(tmp, "set", "20260213-09-00-00")
        os.makedirs(folder)
        os.makedirs(os.path.join(tmp, "set", "masks"))
        rows = []
        for pid in range(4):
            ts = f"20260213-09-00-{10 + pid:02d}"
            present = pid < 3
            m = _write_frame(folder, pid, ts, present, rng)
            mp = ""
            if present:
                mp = f"masks/{pid}.png"
                cv2.imwrite(os.path.join(tmp, "set", mp), m.astype(np.uint8) * 255)
            rows.append({"ir_path": f"20260213-09-00-00/pig{pid}_ir_vis_{ts}.jpg", "mask_path": mp,
                         "label": "present" if present else "absent", "pig_id": pid,
                         "timestamp_folder": "20260213-09-00-00", "pig_timestamp": ts, "how": "hand"})
        labels = os.path.join(tmp, "set", "labels.csv")
        with open(labels, "w", newline="") as f:
            w = csv.DictWriter(f, fieldnames=list(rows[0]))
            w.writeheader()
            w.writerows(rows)
        model = VulvaUNet(pretrained=False).eval()
        ckpt = {"threshold": 0.5, "min_area": 25, "presence_threshold": 0.0, "inputs": "ir"}
        frames = evaluate(model, ckpt, read_rows(labels), "cpu")
        assert len(frames) == 4 and {f["verdict"] for f in frames} <= {
            "false alarm", "missed", "dropped by cut", "both visible", "both not visible"}
        assert all(f["human_status"] == "ok" for f in frames if f["label"] == "present")
        assert abs(float(frames[0]["human_width_mm"]) - 40) < 8          # hand mask measured on the depth
        report = summarise(frames, 0.0)
        assert report[0].startswith("4 labelled frames: 3 with a visible vulva, 1 without")
        from seg_model import save_checkpoint
        ck = os.path.join(tmp, "m.pth")
        save_checkpoint(model, ck, ckpt)
        out = subprocess.run([PY, os.path.join(ROOT, "vulva", "eval_seg.py"), "--model", ck, "--labels", labels,
                              "--out-dir", os.path.join(tmp, "ev")], capture_output=True, text=True, cwd=tmp)
        assert out.returncode == 0, out.stderr[-2000:]
        assert all(os.path.exists(os.path.join(tmp, "ev", n)) for n in ("eval_report.txt", "eval_frames.csv",
                                                                         "eval_sheet.jpg"))
    finally:
        shutil.rmtree(tmp)


if __name__ == "__main__":
    fails = 0
    for name, fn in sorted((n, f) for n, f in globals().items() if n.startswith("test_") and callable(f)):
        try:
            fn()
            print(f"PASS  {name}")
        except Exception as e:  # noqa: BLE001
            fails += 1
            print(f"FAIL  {name}: {type(e).__name__}: {e}")
    print(f"\n{'All' if not fails else fails} {'checks passed' if not fails else 'failed'}")
    sys.exit(1 if fails else 0)
