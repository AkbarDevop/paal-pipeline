#!/usr/bin/env python3
"""Read Lucas's estrus workbook into one tidy sow-day table.

The workbook (Teams > PAAL Working Group > AI Imaging > Files > Ground Truth,
e.g. Feb_2026.xlsx) has one sheet per stall ("0" ... "19"). Each sheet has a
header row with DAY, DATE, CALIPER, BS, V. W., V. L., V. V., MORNING, EVENING
and one row per day after the sows were moved in (day 0 = MOVED).

  V. W. / V. L.  vulva width / length from calipers, mm
  V. V.          Lucas's "vulva volume" = V.W. x V.L. / 200 (an area index, cm^2 / 2)
  MORNING/EVENING  20 = heat observed at that check

Output columns: pig_id, tag, day, date, bcs_caliper, bcs, vulva_width_mm,
vulva_length_mm, vulva_index, heat_morning, heat_evening, heat, action.

Usage:
    python vulva/truth.py "Feb_2026.xlsx" --out outputs/vulva/sow_day_truth.csv
"""

from __future__ import annotations

import argparse
import csv
import os
import statistics
from datetime import date, datetime, timedelta

HEADERS = {
    "DAY": "day",
    "DATE": "date",
    "CALIPER": "bcs_caliper",
    "BS": "bcs",
    "V. W.": "vulva_width_mm",
    "V. L.": "vulva_length_mm",
    "V. V.": "vulva_index",
    "MORNING": "heat_morning",
    "EVENING": "heat_evening",
}
FIELDS = ["pig_id", "tag", "day", "date", "bcs_caliper", "bcs", "vulva_width_mm",
          "vulva_length_mm", "vulva_index", "heat_morning", "heat_evening", "heat", "action"]
HEAT_CODE = 20


def _num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).strip().replace(",", "."))
    except ValueError:
        return None


def _date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    if isinstance(v, (int, float)):            # Excel serial day
        return date(1899, 12, 30) + timedelta(days=int(v))
    if isinstance(v, str):
        for fmt in ("%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d"):
            try:
                return datetime.strptime(v.strip(), fmt).date()
            except ValueError:
                pass
    return None


def _norm(s):
    return " ".join(str(s).upper().split()) if s is not None else ""


def read_sheet(ws, pig_id):
    """Return the daily rows of one sow sheet."""
    cells = [list(r) for r in ws.iter_rows(values_only=True)]
    tag = ""
    for row in cells[:10]:
        if row and _norm(row[0]) == "NUMBER" and len(row) > 1 and row[1] is not None:
            tag = str(row[1]).split(".")[0]
    cols, header_row = {}, None
    for i, row in enumerate(cells):
        names = [_norm(c) for c in row]
        if "V. W." in names and "V. L." in names:
            header_row = i
            for j, n in enumerate(names):
                if n in HEADERS:
                    cols[HEADERS[n]] = j
            # DAY / DATE sit one row higher in Lucas's layout
            for k in range(max(0, i - 2), i + 1):
                for j, n in enumerate(_norm(c) for c in cells[k]):
                    if n in ("DAY", "DATE") and HEADERS[n] not in cols:
                        cols[HEADERS[n]] = j
            break
    if header_row is None:
        return []
    action_col = None
    for k in range(max(0, header_row - 2), header_row + 1):
        for j, c in enumerate(cells[k]):
            if _norm(c) == "ACTION":
                action_col = j
    out = []
    for row in cells[header_row + 1:]:
        get = lambda key: row[cols[key]] if key in cols and cols[key] < len(row) else None  # noqa: E731
        d = _date(get("date"))
        day = _num(get("day"))
        if d is None or day is None:
            continue
        w, l = _num(get("vulva_width_mm")), _num(get("vulva_length_mm"))
        hm, he = _num(get("heat_morning")), _num(get("heat_evening"))
        action = row[action_col] if action_col is not None and action_col < len(row) else ""
        if w is None and l is None and hm is None and he is None and not action:
            continue
        idx = _num(get("vulva_index"))
        if (idx is None or idx == 0) and w and l:
            idx = w * l / 200.0
        out.append({
            "pig_id": pig_id, "tag": tag, "day": int(day), "date": d.isoformat(),
            "bcs_caliper": _num(get("bcs_caliper")), "bcs": _num(get("bcs")),
            "vulva_width_mm": w, "vulva_length_mm": l,
            "vulva_index": round(idx, 3) if (idx and w and l) else None,
            "heat_morning": int(hm == HEAT_CODE) if hm is not None else None,
            "heat_evening": int(he == HEAT_CODE) if he is not None else None,
            "heat": int(HEAT_CODE in (hm, he)),
            "action": str(action or "").strip(),
        })
    return out


def read_workbook(path, num_stalls=20):
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True, read_only=True)
    rows = []
    for pid in range(num_stalls):
        if str(pid) in wb.sheetnames:
            rows.extend(read_sheet(wb[str(pid)], pid))
    return rows


def heat_response(rows):
    """Per sow: baseline (days before first heat, at most days 1-2) vs the peak
    from the day before heat to the last heat day. Also day-to-day caliper noise
    after heat, which is the noise floor any camera measurement is compared to."""
    by = {}
    for r in rows:
        by.setdefault(r["pig_id"], []).append(r)
    summary, diffs = [], {"vulva_width_mm": [], "vulva_length_mm": [], "vulva_index": []}
    for pid, rs in sorted(by.items()):
        rs = sorted((r for r in rs if r["day"] >= 1 and r["vulva_width_mm"]), key=lambda r: r["day"])
        heat_days = [r["day"] for r in rs if r["heat"]]
        rec = {"pig_id": pid, "heat_days": "|".join(map(str, heat_days))}
        if heat_days:
            first, last = min(heat_days), max(heat_days)
            base = [r for r in rs if r["day"] < first and r["day"] <= 2] or rs[:1]
            win = [r for r in rs if first - 1 <= r["day"] <= last]
            for k in ("vulva_width_mm", "vulva_length_mm", "vulva_index"):
                b = statistics.mean(r[k] for r in base)
                p = max(r[k] for r in win)
                rec[f"{k}_baseline"] = round(b, 2)
                rec[f"{k}_peak"] = round(p, 2)
                rec[f"{k}_rise_pct"] = round(100 * (p - b) / b, 1)
            post = [r for r in rs if r["day"] > last + 1]
            for a, b in zip(post, post[1:]):
                if b["day"] == a["day"] + 1:
                    for k in diffs:
                        diffs[k].append(b[k] - a[k])
        summary.append(rec)
    noise = {k: round(statistics.stdev(v), 2) if len(v) > 1 else None for k, v in diffs.items()}
    return summary, noise


def write_csv(rows, path, fields):
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    with open(path, "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=fields, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook")
    ap.add_argument("--out", default=os.path.join("outputs", "vulva", "sow_day_truth.csv"))
    ap.add_argument("--stalls", type=int, default=20)
    args = ap.parse_args()

    rows = read_workbook(args.workbook, args.stalls)
    write_csv(rows, args.out, FIELDS)
    summary, noise = heat_response(rows)
    sum_path = os.path.splitext(args.out)[0] + "_heat_response.csv"
    sum_fields = ["pig_id", "heat_days"] + [f"{k}_{s}" for k in ("vulva_width_mm", "vulva_length_mm", "vulva_index")
                                           for s in ("baseline", "peak", "rise_pct")]
    write_csv(summary, sum_path, sum_fields)

    sows = sorted({r["pig_id"] for r in rows})
    in_heat = [s for s in summary if s["heat_days"]]
    rises = sorted(s["vulva_index_rise_pct"] for s in in_heat)
    print(f"Read {len(rows)} sow-days for {len(sows)} sows -> {args.out}")
    print(f"Sows with heat recorded: {len(in_heat)} of {len(sows)}")
    if rises:
        med = statistics.median(rises)
        print(f"Vulva index rise from baseline to heat peak: median {med:.1f}% "
              f"(>10% in {sum(r > 10 for r in rises)} of {len(rises)}, >5% in {sum(r > 5 for r in rises)})")
        for k, label in (("vulva_width_mm", "width"), ("vulva_length_mm", "length")):
            print(f"  {label}: median rise {statistics.median(s[k + '_rise_pct'] for s in in_heat):.1f}%")
    print("Caliper day-to-day noise after heat (SD of daily change): "
          + ", ".join(f"{k} {v}" for k, v in noise.items()))
    print(f"Per-sow table -> {sum_path}")


if __name__ == "__main__":
    main()
